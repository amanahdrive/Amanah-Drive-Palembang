from fastapi import FastAPI, APIRouter, HTTPException, UploadFile, File
from fastapi.responses import HTMLResponse, Response
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import io
import logging
from pathlib import Path
from pydantic import BaseModel, Field
from typing import List, Optional, Literal
import uuid
from datetime import datetime, timezone, timedelta
from openpyxl import Workbook, load_workbook


ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]

app = FastAPI()
api_router = APIRouter(prefix="/api")


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


# ---------- Models ----------
StudentStatus = Literal["aktif", "on_progress", "selesai"]
TransactionType = Literal["masuk", "keluar"]


class Package(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    nama: str
    harga: int
    total_pertemuan: int
    deskripsi: str = ""


class StudentCreate(BaseModel):
    nama: str
    telepon: str
    paket: str
    instruktur: Optional[str] = ""
    total_pertemuan: int = 10
    pertemuan_selesai: int = 0
    tanggal_daftar: Optional[str] = None
    status: StudentStatus = "aktif"
    harga_paket: int = 0
    sudah_dibayar: int = 0
    catatan: Optional[str] = ""


class Student(StudentCreate):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    created_at: str = Field(default_factory=now_iso)
    updated_at: str = Field(default_factory=now_iso)


class StudentUpdate(BaseModel):
    nama: Optional[str] = None
    telepon: Optional[str] = None
    paket: Optional[str] = None
    instruktur: Optional[str] = None
    total_pertemuan: Optional[int] = None
    pertemuan_selesai: Optional[int] = None
    tanggal_daftar: Optional[str] = None
    status: Optional[StudentStatus] = None
    harga_paket: Optional[int] = None
    sudah_dibayar: Optional[int] = None
    catatan: Optional[str] = None


class TransactionCreate(BaseModel):
    jenis: TransactionType
    kategori: str
    jumlah: int
    keterangan: Optional[str] = ""
    tanggal: Optional[str] = None
    siswa_id: Optional[str] = None
    siswa_nama: Optional[str] = None


class Transaction(TransactionCreate):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    created_at: str = Field(default_factory=now_iso)


class TransactionUpdate(BaseModel):
    jenis: Optional[TransactionType] = None
    kategori: Optional[str] = None
    jumlah: Optional[int] = None
    keterangan: Optional[str] = None
    tanggal: Optional[str] = None
    siswa_id: Optional[str] = None
    siswa_nama: Optional[str] = None


class ScheduleCreate(BaseModel):
    tanggal: str  # YYYY-MM-DD
    jam_mulai: str  # HH:MM
    jam_selesai: str  # HH:MM
    instruktur: str
    siswa_id: Optional[str] = None
    siswa_nama: Optional[str] = None
    catatan: Optional[str] = ""


class Schedule(ScheduleCreate):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    created_at: str = Field(default_factory=now_iso)


class ScheduleUpdate(BaseModel):
    tanggal: Optional[str] = None
    jam_mulai: Optional[str] = None
    jam_selesai: Optional[str] = None
    instruktur: Optional[str] = None
    siswa_id: Optional[str] = None
    siswa_nama: Optional[str] = None
    catatan: Optional[str] = None


class PaymentReminder(BaseModel):
    siswa_id: str
    nama: str
    telepon: str
    paket: str
    harga_paket: int
    sudah_dibayar: int
    sisa_bayar: int
    status: str


class DashboardSummary(BaseModel):
    total_siswa: int
    siswa_aktif: int
    siswa_on_progress: int
    siswa_selesai: int
    kas_masuk_bulan: int
    kas_keluar_bulan: int
    sisa_kas: int
    kas_masuk_total: int
    kas_keluar_total: int
    chart_7hari: List[dict]
    pengingat_pembayaran: List[PaymentReminder]
    jadwal_hari_ini: List[dict]


DEFAULT_PACKAGES = [
    {"id": "paket-mobil-manual", "nama": "Mobil Manual", "harga": 1500000, "total_pertemuan": 10,
     "deskripsi": "Paket kursus mobil manual 10x pertemuan"},
    {"id": "paket-mobil-matic", "nama": "Mobil Matic", "harga": 1400000, "total_pertemuan": 8,
     "deskripsi": "Paket kursus mobil matic 8x pertemuan"},
    {"id": "paket-motor", "nama": "Motor", "harga": 500000, "total_pertemuan": 6,
     "deskripsi": "Paket kursus motor 6x pertemuan"},
]

INCOME_CATEGORIES = ["Pembayaran Kursus", "Pendaftaran", "Lainnya"]
EXPENSE_CATEGORIES = ["Gaji Instruktur", "BBM", "Maintenance Kendaraan", "Operasional", "Lainnya"]


# ---------- Startup: seed packages ----------
@app.on_event("startup")
async def seed_packages():
    for p in DEFAULT_PACKAGES:
        existing = await db.packages.find_one({"id": p["id"]})
        if not existing:
            await db.packages.insert_one(p.copy())


# ---------- Helpers ----------
async def recompute_student_payment(student_id: str):
    """Sum semua transaksi 'masuk' dengan siswa_id, update sudah_dibayar."""
    cursor = db.transactions.find(
        {"siswa_id": student_id, "jenis": "masuk", "deleted_at": {"$exists": False}},
        {"_id": 0, "jumlah": 1},
    )
    total = 0
    async for row in cursor:
        total += row.get("jumlah", 0)
    await db.students.update_one({"id": student_id}, {"$set": {"sudah_dibayar": total, "updated_at": now_iso()}})


def rupiah(value: int) -> str:
    sign = "-" if value < 0 else ""
    v = abs(int(value))
    s = f"{v:,}".replace(",", ".")
    return f"{sign}Rp {s}"


def parse_date(iso: str) -> Optional[datetime]:
    try:
        if not iso:
            return None
        d = datetime.fromisoformat(iso.replace("Z", "+00:00"))
        if d.tzinfo is None:
            d = d.replace(tzinfo=timezone.utc)
        return d
    except Exception:
        return None


# ---------- Routes ----------
@api_router.get("/")
async def root():
    return {"message": "Amanah Drive API", "version": "1.1"}


# Packages
@api_router.get("/packages", response_model=List[Package])
async def list_packages():
    rows = await db.packages.find({}, {"_id": 0}).to_list(100)
    return [Package(**r) for r in rows]


# Categories
@api_router.get("/categories")
async def list_categories():
    return {"masuk": INCOME_CATEGORIES, "keluar": EXPENSE_CATEGORIES}


# Students
@api_router.get("/students", response_model=List[Student])
async def list_students(status: Optional[str] = None):
    query = {"deleted_at": {"$exists": False}}
    if status and status != "semua":
        query["status"] = status
    rows = await db.students.find(query, {"_id": 0}).sort("created_at", -1).to_list(1000)
    return [Student(**r) for r in rows]


@api_router.post("/students", response_model=Student)
async def create_student(payload: StudentCreate):
    data = payload.dict()
    if not data.get("tanggal_daftar"):
        data["tanggal_daftar"] = now_iso()
    student = Student(**data)
    await db.students.insert_one(student.dict().copy())
    return student


@api_router.get("/students/{student_id}", response_model=Student)
async def get_student(student_id: str):
    row = await db.students.find_one({"id": student_id, "deleted_at": {"$exists": False}}, {"_id": 0})
    if not row:
        raise HTTPException(status_code=404, detail="Student not found")
    return Student(**row)


@api_router.put("/students/{student_id}", response_model=Student)
async def update_student(student_id: str, payload: StudentUpdate):
    updates = {k: v for k, v in payload.dict().items() if v is not None}
    if not updates:
        raise HTTPException(status_code=400, detail="No fields to update")
    updates["updated_at"] = now_iso()
    if "pertemuan_selesai" in updates or "total_pertemuan" in updates:
        existing = await db.students.find_one({"id": student_id}, {"_id": 0})
        if existing:
            selesai = updates.get("pertemuan_selesai", existing.get("pertemuan_selesai", 0))
            total = updates.get("total_pertemuan", existing.get("total_pertemuan", 0))
            if selesai >= total and total > 0:
                updates["status"] = "selesai"
            elif selesai > 0:
                updates["status"] = "on_progress"
    result = await db.students.update_one({"id": student_id}, {"$set": updates})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Student not found")
    row = await db.students.find_one({"id": student_id}, {"_id": 0})
    return Student(**row)


@api_router.delete("/students/{student_id}")
async def delete_student(student_id: str):
    result = await db.students.update_one(
        {"id": student_id},
        {"$set": {"deleted_at": now_iso()}},
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Student not found")
    return {"ok": True}


# Transactions
@api_router.get("/transactions", response_model=List[Transaction])
async def list_transactions(jenis: Optional[str] = None, bulan_ini: Optional[bool] = False):
    query = {"deleted_at": {"$exists": False}}
    if jenis and jenis != "semua":
        query["jenis"] = jenis
    if bulan_ini:
        now = datetime.now(timezone.utc)
        start = datetime(now.year, now.month, 1, tzinfo=timezone.utc).isoformat()
        query["tanggal"] = {"$gte": start}
    rows = await db.transactions.find(query, {"_id": 0}).sort("tanggal", -1).to_list(2000)
    return [Transaction(**r) for r in rows]


@api_router.post("/transactions", response_model=Transaction)
async def create_transaction(payload: TransactionCreate):
    data = payload.dict()
    if not data.get("tanggal"):
        data["tanggal"] = now_iso()
    tx = Transaction(**data)
    await db.transactions.insert_one(tx.dict())
    if tx.jenis == "masuk" and tx.siswa_id:
        await recompute_student_payment(tx.siswa_id)
    return tx


@api_router.get("/transactions/{tx_id}", response_model=Transaction)
async def get_transaction(tx_id: str):
    row = await db.transactions.find_one({"id": tx_id, "deleted_at": {"$exists": False}}, {"_id": 0})
    if not row:
        raise HTTPException(status_code=404, detail="Transaction not found")
    return Transaction(**row)


@api_router.put("/transactions/{tx_id}", response_model=Transaction)
async def update_transaction(tx_id: str, payload: TransactionUpdate):
    updates = {k: v for k, v in payload.dict().items() if v is not None}
    if not updates:
        raise HTTPException(status_code=400, detail="No fields to update")
    before = await db.transactions.find_one({"id": tx_id}, {"_id": 0})
    result = await db.transactions.update_one({"id": tx_id}, {"$set": updates})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Transaction not found")
    row = await db.transactions.find_one({"id": tx_id}, {"_id": 0})
    # Recompute siswa payments on both before/after siswa_id
    for sid in {(before or {}).get("siswa_id"), row.get("siswa_id")}:
        if sid:
            await recompute_student_payment(sid)
    return Transaction(**row)


@api_router.delete("/transactions/{tx_id}")
async def delete_transaction(tx_id: str):
    before = await db.transactions.find_one({"id": tx_id}, {"_id": 0})
    result = await db.transactions.update_one(
        {"id": tx_id},
        {"$set": {"deleted_at": now_iso()}},
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Transaction not found")
    if before and before.get("siswa_id"):
        await recompute_student_payment(before["siswa_id"])
    return {"ok": True}


# Schedules
@api_router.get("/schedules", response_model=List[Schedule])
async def list_schedules(tanggal: Optional[str] = None, instruktur: Optional[str] = None):
    query = {"deleted_at": {"$exists": False}}
    if tanggal:
        query["tanggal"] = tanggal
    if instruktur:
        query["instruktur"] = instruktur
    rows = await db.schedules.find(query, {"_id": 0}).sort([("tanggal", 1), ("jam_mulai", 1)]).to_list(2000)
    return [Schedule(**r) for r in rows]


@api_router.post("/schedules", response_model=Schedule)
async def create_schedule(payload: ScheduleCreate):
    sch = Schedule(**payload.dict())
    await db.schedules.insert_one(sch.dict())
    return sch


@api_router.put("/schedules/{sch_id}", response_model=Schedule)
async def update_schedule(sch_id: str, payload: ScheduleUpdate):
    updates = {k: v for k, v in payload.dict().items() if v is not None}
    if not updates:
        raise HTTPException(status_code=400, detail="No fields to update")
    result = await db.schedules.update_one({"id": sch_id}, {"$set": updates})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Schedule not found")
    row = await db.schedules.find_one({"id": sch_id}, {"_id": 0})
    return Schedule(**row)


@api_router.delete("/schedules/{sch_id}")
async def delete_schedule(sch_id: str):
    result = await db.schedules.update_one(
        {"id": sch_id},
        {"$set": {"deleted_at": now_iso()}},
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Schedule not found")
    return {"ok": True}


@api_router.get("/instrukturs", response_model=List[str])
async def list_instrukturs():
    """Daftar unik instruktur yang pernah dicatat (dari siswa + jadwal)."""
    s = await db.students.distinct("instruktur", {"deleted_at": {"$exists": False}})
    j = await db.schedules.distinct("instruktur", {"deleted_at": {"$exists": False}})
    merged = sorted({x for x in list(s) + list(j) if x})
    return merged


# Dashboard
@api_router.get("/dashboard", response_model=DashboardSummary)
async def dashboard():
    students = await db.students.find({"deleted_at": {"$exists": False}}, {"_id": 0}).to_list(5000)
    txs = await db.transactions.find({"deleted_at": {"$exists": False}}, {"_id": 0}).to_list(10000)

    total_siswa = len(students)
    siswa_aktif = sum(1 for s in students if s.get("status") == "aktif")
    siswa_on_progress = sum(1 for s in students if s.get("status") == "on_progress")
    siswa_selesai = sum(1 for s in students if s.get("status") == "selesai")

    kas_masuk_total = sum(t.get("jumlah", 0) for t in txs if t.get("jenis") == "masuk")
    kas_keluar_total = sum(t.get("jumlah", 0) for t in txs if t.get("jenis") == "keluar")
    sisa_kas = kas_masuk_total - kas_keluar_total

    now = datetime.now(timezone.utc)
    start_month = datetime(now.year, now.month, 1, tzinfo=timezone.utc)
    kas_masuk_bulan = 0
    kas_keluar_bulan = 0
    for t in txs:
        tgl = parse_date(t.get("tanggal", ""))
        if tgl and tgl >= start_month:
            if t.get("jenis") == "masuk":
                kas_masuk_bulan += t.get("jumlah", 0)
            else:
                kas_keluar_bulan += t.get("jumlah", 0)

    chart = []
    for i in range(6, -1, -1):
        day = (now - timedelta(days=i)).replace(hour=0, minute=0, second=0, microsecond=0)
        next_day = day + timedelta(days=1)
        masuk = 0
        keluar = 0
        for t in txs:
            tgl = parse_date(t.get("tanggal", ""))
            if tgl and day <= tgl < next_day:
                if t.get("jenis") == "masuk":
                    masuk += t.get("jumlah", 0)
                else:
                    keluar += t.get("jumlah", 0)
        chart.append({"tanggal": day.isoformat(), "label": day.strftime("%a"), "masuk": masuk, "keluar": keluar})

    # Payment reminders
    reminders: List[PaymentReminder] = []
    for s in students:
        if s.get("status") == "selesai":
            continue
        harga = int(s.get("harga_paket") or 0)
        dibayar = int(s.get("sudah_dibayar") or 0)
        sisa = harga - dibayar
        if sisa > 0 and harga > 0:
            reminders.append(PaymentReminder(
                siswa_id=s["id"],
                nama=s.get("nama", ""),
                telepon=s.get("telepon", ""),
                paket=s.get("paket", ""),
                harga_paket=harga,
                sudah_dibayar=dibayar,
                sisa_bayar=sisa,
                status=s.get("status", "aktif"),
            ))
    reminders.sort(key=lambda r: r.sisa_bayar, reverse=True)

    # Jadwal hari ini
    today_str = now.strftime("%Y-%m-%d")
    jadwal_today = await db.schedules.find(
        {"tanggal": today_str, "deleted_at": {"$exists": False}},
        {"_id": 0},
    ).sort("jam_mulai", 1).to_list(100)

    return DashboardSummary(
        total_siswa=total_siswa,
        siswa_aktif=siswa_aktif,
        siswa_on_progress=siswa_on_progress,
        siswa_selesai=siswa_selesai,
        kas_masuk_bulan=kas_masuk_bulan,
        kas_keluar_bulan=kas_keluar_bulan,
        sisa_kas=sisa_kas,
        kas_masuk_total=kas_masuk_total,
        kas_keluar_total=kas_keluar_total,
        chart_7hari=chart,
        pengingat_pembayaran=reminders[:10],
        jadwal_hari_ini=jadwal_today,
    )


# Reports (HTML for PDF printing)
@api_router.get("/reports/cashflow", response_class=HTMLResponse)
async def report_cashflow(bulan: Optional[str] = None):
    """Return HTML laporan cashflow bulanan (siap print ke PDF).
    bulan format: YYYY-MM. default: bulan berjalan.
    """
    now = datetime.now(timezone.utc)
    if bulan:
        try:
            year, month = [int(x) for x in bulan.split("-")]
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid bulan format, use YYYY-MM")
    else:
        year, month = now.year, now.month

    start = datetime(year, month, 1, tzinfo=timezone.utc)
    end = datetime(year + 1, 1, 1, tzinfo=timezone.utc) if month == 12 else datetime(year, month + 1, 1, tzinfo=timezone.utc)

    rows = await db.transactions.find(
        {"deleted_at": {"$exists": False}},
        {"_id": 0},
    ).sort("tanggal", 1).to_list(5000)
    filtered = [t for t in rows if (parse_date(t.get("tanggal", "")) or start) >= start and (parse_date(t.get("tanggal", "")) or end) < end]

    masuk_total = sum(t["jumlah"] for t in filtered if t["jenis"] == "masuk")
    keluar_total = sum(t["jumlah"] for t in filtered if t["jenis"] == "keluar")
    sisa = masuk_total - keluar_total

    bulan_nama = ["Januari","Februari","Maret","April","Mei","Juni","Juli","Agustus","September","Oktober","November","Desember"][month - 1]

    tx_rows_html = []
    for t in filtered:
        tgl = parse_date(t.get("tanggal", ""))
        tgl_str = tgl.strftime("%d/%m/%Y %H:%M") if tgl else "-"
        is_masuk = t.get("jenis") == "masuk"
        amount_color = "#10B981" if is_masuk else "#EF4444"
        sign = "+" if is_masuk else "-"
        tx_rows_html.append(f"""
          <tr>
            <td>{tgl_str}</td>
            <td>{"Kas Masuk" if is_masuk else "Kas Keluar"}</td>
            <td>{t.get("kategori", "")}</td>
            <td>{(t.get("keterangan") or "-")}</td>
            <td>{(t.get("siswa_nama") or "-")}</td>
            <td style="color:{amount_color};text-align:right;font-weight:700;">{sign}{rupiah(t.get("jumlah", 0))}</td>
          </tr>
        """)

    html = f"""<!doctype html>
<html lang="id">
<head>
<meta charset="utf-8"/>
<title>Laporan Cashflow {bulan_nama} {year} - Amanah Drive</title>
<style>
  * {{ box-sizing: border-box; }}
  body {{ font-family: -apple-system, Segoe UI, Roboto, Arial, sans-serif; color: #111827; margin: 24px; }}
  .header {{ display:flex; align-items:center; gap:16px; border-bottom: 3px solid #00267A; padding-bottom: 12px; }}
  .brand {{ font-size: 24px; font-weight: 900; color: #00267A; }}
  .tagline {{ font-size: 12px; color: #6B7280; }}
  h1 {{ font-size: 18px; margin: 16px 0 4px; color: #00267A; }}
  .period {{ font-size: 13px; color: #374151; margin-bottom: 16px; }}
  .summary {{ display:flex; gap: 12px; margin: 16px 0; }}
  .card {{ flex:1; border: 1px solid #E5E7EB; border-radius: 10px; padding: 12px 14px; }}
  .card .label {{ font-size: 11px; color: #6B7280; font-weight:600; }}
  .card .value {{ font-size: 18px; font-weight: 900; margin-top: 4px; }}
  .masuk .value {{ color: #10B981; }}
  .keluar .value {{ color: #EF4444; }}
  .sisa .value {{ color: #00267A; }}
  table {{ width:100%; border-collapse: collapse; margin-top: 8px; font-size: 12px; }}
  th {{ background: #00267A; color: #fff; padding: 8px; text-align: left; }}
  td {{ padding: 8px; border-bottom: 1px solid #E5E7EB; }}
  tr:nth-child(even) td {{ background: #F9FAFB; }}
  .foot {{ margin-top: 24px; font-size: 11px; color: #6B7280; text-align: center; }}
</style>
</head>
<body>
  <div class="header">
    <div>
      <div class="brand">Amanah Drive</div>
      <div class="tagline">Terampil Mengemudi, Aman di Jalan.</div>
    </div>
  </div>
  <h1>Laporan Cashflow</h1>
  <div class="period">Periode: <b>{bulan_nama} {year}</b> &nbsp;|&nbsp; Dicetak: {now.strftime("%d/%m/%Y %H:%M UTC")}</div>
  <div class="summary">
    <div class="card masuk"><div class="label">TOTAL KAS MASUK</div><div class="value">{rupiah(masuk_total)}</div></div>
    <div class="card keluar"><div class="label">TOTAL KAS KELUAR</div><div class="value">{rupiah(keluar_total)}</div></div>
    <div class="card sisa"><div class="label">SISA KAS</div><div class="value">{rupiah(sisa)}</div></div>
  </div>
  <table>
    <thead>
      <tr>
        <th>Tanggal</th>
        <th>Jenis</th>
        <th>Kategori</th>
        <th>Keterangan</th>
        <th>Siswa</th>
        <th style="text-align:right;">Jumlah</th>
      </tr>
    </thead>
    <tbody>
      {''.join(tx_rows_html) if tx_rows_html else '<tr><td colspan="6" style="text-align:center;padding:16px;color:#6B7280;">Belum ada transaksi di bulan ini.</td></tr>'}
    </tbody>
  </table>
  <div class="foot">© {year} Amanah Drive — amanahdrive_plg</div>
</body>
</html>
"""
    return HTMLResponse(content=html)


# ---------- Import Siswa Excel ----------
STUDENT_IMPORT_COLUMNS = [
    "nama", "telepon", "paket", "instruktur",
    "total_pertemuan", "pertemuan_selesai",
    "harga_paket", "sudah_dibayar", "status", "catatan",
]


@api_router.get("/import-siswa/template")
async def download_import_template():
    """Return an XLSX template file users can fill and upload."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Siswa"
    ws.append(STUDENT_IMPORT_COLUMNS)
    # Example rows
    ws.append(["Budi Santoso", "081234567890", "Mobil Manual", "Pak Agus", 10, 2, 1500000, 500000, "on_progress", "Contoh baris"])
    ws.append(["Siti Rahma", "081298765432", "Mobil Matic", "Bu Dewi", 8, 0, 1400000, 0, "aktif", ""])
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template-impor-siswa.xlsx"},
    )


@api_router.post("/import-siswa")
async def import_students(file: UploadFile = File(...)):
    """Import students from an uploaded XLSX file.
    Expected columns (header row): nama, telepon, paket, instruktur,
    total_pertemuan, pertemuan_selesai, harga_paket, sudah_dibayar, status, catatan.
    Rows missing nama or telepon are skipped.
    """
    try:
        raw = await file.read()
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="File XLSX tidak valid")

    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        raise HTTPException(status_code=400, detail="File kosong")
    header_norm = [str(h or "").strip().lower() for h in header]
    index = {col: (header_norm.index(col) if col in header_norm else -1) for col in STUDENT_IMPORT_COLUMNS}

    created = 0
    skipped = 0
    errors: List[str] = []

    def get(row, col):
        i = index.get(col, -1)
        if i < 0 or i >= len(row):
            return None
        return row[i]

    def to_int(v, default=0):
        try:
            if v is None or v == "":
                return default
            return int(float(v))
        except Exception:
            return default

    for row_no, row in enumerate(rows, start=2):
        if row is None:
            continue
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
            continue
        nama = (get(row, "nama") or "").strip() if isinstance(get(row, "nama"), str) else str(get(row, "nama") or "").strip()
        telepon = (get(row, "telepon") or "")
        telepon = str(telepon).strip()
        paket = str(get(row, "paket") or "").strip()
        if not nama or not telepon:
            skipped += 1
            errors.append(f"Baris {row_no}: nama/telepon kosong")
            continue
        total = to_int(get(row, "total_pertemuan"), 10)
        selesai = to_int(get(row, "pertemuan_selesai"), 0)
        harga = to_int(get(row, "harga_paket"), 0)
        dibayar = to_int(get(row, "sudah_dibayar"), 0)
        status_val = str(get(row, "status") or "aktif").strip().lower()
        if status_val not in ("aktif", "on_progress", "selesai"):
            status_val = "aktif"
        # Auto-adjust status from pertemuan
        if total > 0 and selesai >= total:
            status_val = "selesai"
        elif selesai > 0 and status_val == "aktif":
            status_val = "on_progress"
        payload = StudentCreate(
            nama=nama,
            telepon=telepon,
            paket=paket or "Mobil Manual",
            instruktur=str(get(row, "instruktur") or "").strip(),
            total_pertemuan=total,
            pertemuan_selesai=selesai,
            harga_paket=harga,
            sudah_dibayar=dibayar,
            status=status_val,  # type: ignore[arg-type]
            catatan=str(get(row, "catatan") or "").strip(),
        )
        data = payload.dict()
        data["tanggal_daftar"] = now_iso()
        student = Student(**data)
        await db.students.insert_one(student.dict().copy())
        created += 1

    return {"created": created, "skipped": skipped, "errors": errors[:20]}


# ---------- Rekap Instruktur ----------
class RekapInstruktur(BaseModel):
    instruktur: str
    total_sesi: int
    total_jam: float
    estimasi_bagi_hasil: int


class RekapInstrukturResponse(BaseModel):
    bulan: str
    rate_per_jam: int
    items: List[RekapInstruktur]
    total_sesi: int
    total_jam: float
    total_bagi_hasil: int


@api_router.get("/reports/instrukturs", response_model=RekapInstrukturResponse)
async def report_instrukturs(bulan: Optional[str] = None, rate_per_jam: int = 50000):
    """Rekap jam mengajar & bagi hasil per instruktur dari jadwal bulan tertentu."""
    now = datetime.now(timezone.utc)
    if bulan:
        try:
            year, month = [int(x) for x in bulan.split("-")]
        except Exception:
            raise HTTPException(status_code=400, detail="Invalid bulan format, use YYYY-MM")
    else:
        year, month = now.year, now.month
    prefix = f"{year:04d}-{month:02d}-"

    rows = await db.schedules.find(
        {"deleted_at": {"$exists": False}, "tanggal": {"$regex": f"^{prefix}"}},
        {"_id": 0},
    ).to_list(5000)

    def hours_between(a: str, b: str) -> float:
        try:
            ha, ma = [int(x) for x in a.split(":")]
            hb, mb = [int(x) for x in b.split(":")]
            return max(0.0, (hb + mb / 60.0) - (ha + ma / 60.0))
        except Exception:
            return 0.0

    agg: dict[str, dict] = {}
    for s in rows:
        key = (s.get("instruktur") or "-").strip() or "-"
        d = agg.setdefault(key, {"total_sesi": 0, "total_jam": 0.0})
        d["total_sesi"] += 1
        d["total_jam"] += hours_between(s.get("jam_mulai", ""), s.get("jam_selesai", ""))

    items = []
    total_sesi = 0
    total_jam = 0.0
    for name, d in sorted(agg.items(), key=lambda kv: kv[1]["total_jam"], reverse=True):
        hrs = round(d["total_jam"], 2)
        bagi = int(round(hrs * rate_per_jam))
        items.append(RekapInstruktur(instruktur=name, total_sesi=d["total_sesi"], total_jam=hrs, estimasi_bagi_hasil=bagi))
        total_sesi += d["total_sesi"]
        total_jam += hrs

    return RekapInstrukturResponse(
        bulan=f"{year:04d}-{month:02d}",
        rate_per_jam=rate_per_jam,
        items=items,
        total_sesi=total_sesi,
        total_jam=round(total_jam, 2),
        total_bagi_hasil=sum(x.estimasi_bagi_hasil for x in items),
    )


# ---------- Sertifikat Siswa ----------
@api_router.get("/students/{student_id}/sertifikat", response_class=HTMLResponse)
async def student_sertifikat(student_id: str):
    row = await db.students.find_one({"id": student_id, "deleted_at": {"$exists": False}}, {"_id": 0})
    if not row:
        raise HTTPException(status_code=404, detail="Student not found")
    s = Student(**row)
    today = datetime.now(timezone.utc)
    today_id = today.strftime("%d %B %Y")
    bulan_id = ["Januari","Februari","Maret","April","Mei","Juni","Juli","Agustus","September","Oktober","November","Desember"][today.month - 1]
    tgl_str = f"{today.day} {bulan_id} {today.year}"
    cert_no = f"AD/{today.year}/{s.id[:6].upper()}"

    html = f"""<!doctype html>
<html lang=\"id\">
<head>
<meta charset=\"utf-8\"/>
<title>Sertifikat {s.nama} - Amanah Drive</title>
<style>
  @page {{ size: A4 landscape; margin: 0; }}
  * {{ box-sizing: border-box; }}
  body {{
    margin: 0;
    font-family: Georgia, 'Times New Roman', serif;
    background: #f1f5f9;
  }}
  .page {{
    width: 100%;
    min-height: 100vh;
    padding: 32px;
    display: flex;
    align-items: center;
    justify-content: center;
  }}
  .cert {{
    width: 1120px;
    aspect-ratio: 1120/780;
    background: linear-gradient(135deg, #ffffff 0%, #eaf4fb 100%);
    border: 14px double #00267A;
    border-radius: 20px;
    padding: 56px 72px;
    position: relative;
    box-shadow: 0 20px 60px rgba(0,38,122,0.25);
  }}
  .corner {{
    position: absolute; width: 90px; height: 90px;
    border: 4px solid #3DD0F2; border-radius: 10px;
  }}
  .c1 {{ top: 24px; left: 24px; border-right: 0; border-bottom: 0; }}
  .c2 {{ top: 24px; right: 24px; border-left: 0; border-bottom: 0; }}
  .c3 {{ bottom: 24px; left: 24px; border-right: 0; border-top: 0; }}
  .c4 {{ bottom: 24px; right: 24px; border-left: 0; border-top: 0; }}
  .brand {{ text-align: center; }}
  .brand .name {{ font-family: Arial, sans-serif; font-size: 36px; color: #00267A; font-weight: 900; letter-spacing: 1px; }}
  .brand .tag {{ font-family: Arial, sans-serif; font-size: 13px; color: #005b93; letter-spacing: 2px; margin-top: 4px; }}
  .title {{ margin-top: 32px; text-align: center; font-size: 48px; font-weight: 900; color: #00267A; letter-spacing: 3px; }}
  .subtitle {{ text-align: center; font-size: 14px; letter-spacing: 4px; color: #005b93; margin-top: 4px; }}
  .awarded {{ text-align: center; margin-top: 28px; font-size: 15px; color: #334155; }}
  .student {{ text-align: center; margin-top: 10px; font-size: 44px; font-family: 'Brush Script MT', 'Pacifico', cursive; color: #005b93; border-bottom: 2px solid #3DD0F2; padding-bottom: 10px; max-width: 760px; margin-left: auto; margin-right: auto; }}
  .body {{ text-align: center; margin-top: 24px; font-size: 15px; color: #334155; line-height: 1.6; max-width: 820px; margin-left: auto; margin-right: auto; }}
  .body b {{ color: #00267A; }}
  .meta {{ display: flex; justify-content: space-between; align-items: flex-end; margin-top: 48px; padding: 0 60px; }}
  .meta .block {{ text-align: center; min-width: 180px; }}
  .meta .label {{ font-size: 11px; color: #64748b; letter-spacing: 1px; }}
  .meta .value {{ font-size: 14px; font-weight: 800; color: #00267A; margin-top: 2px; }}
  .seal {{ width: 110px; height: 110px; border: 4px solid #00267A; border-radius: 55px; display: inline-flex; align-items: center; justify-content: center; color: #00267A; font-weight: 900; font-family: Arial, sans-serif; font-size: 11px; text-align: center; }}
  .footer {{ text-align: center; margin-top: 22px; font-size: 11px; color: #64748b; letter-spacing: 2px; }}
</style>
</head>
<body>
<div class=\"page\">
  <div class=\"cert\">
    <div class=\"corner c1\"></div><div class=\"corner c2\"></div>
    <div class=\"corner c3\"></div><div class=\"corner c4\"></div>
    <div class=\"brand\">
      <div class=\"name\">AMANAH DRIVE</div>
      <div class=\"tag\">TERAMPIL MENGEMUDI · AMAN DI JALAN</div>
    </div>
    <div class=\"title\">SERTIFIKAT</div>
    <div class=\"subtitle\">KELULUSAN KURSUS MENGEMUDI</div>
    <div class=\"awarded\">Diberikan dengan bangga kepada:</div>
    <div class=\"student\">{s.nama}</div>
    <div class=\"body\">
      Yang telah menyelesaikan <b>Paket {s.paket}</b> sebanyak
      <b>{s.total_pertemuan} pertemuan</b> di Amanah Drive dengan bimbingan instruktur
      <b>{s.instruktur or '-'}</b>, dan dinyatakan <b>LULUS</b> serta memiliki keterampilan
      mengemudi yang baik, aman, dan bertanggung jawab di jalan raya.
    </div>
    <div class=\"meta\">
      <div class=\"block\">
        <div class=\"label\">NOMOR SERTIFIKAT</div>
        <div class=\"value\">{cert_no}</div>
      </div>
      <div class=\"block\">
        <div class=\"seal\">AMANAH<br/>DRIVE<br/>PLG</div>
      </div>
      <div class=\"block\">
        <div class=\"label\">TANGGAL</div>
        <div class=\"value\">{tgl_str}</div>
      </div>
    </div>
    <div class=\"footer\">AMANAH DRIVE PALEMBANG · amanahdrive_plg</div>
  </div>
</div>
</body>
</html>
"""
    return HTMLResponse(content=html)


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
)
logger = logging.getLogger(__name__)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
